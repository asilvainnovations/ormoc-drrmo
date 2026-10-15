#!/usr/bin/env python3
"""Generate reference data for the platform from the City's source-of-truth workbooks.

Inputs  (reference/):
  Ormoc_CDRRMO-EOC_Institutional_Assessment_Tool_v1.1.xlsx          -> parts, domains, 140 indicators, crosswalk
  Ormoc_CDRRMO_EOC_Master_Requirements_Traceability_and_140_...xlsx -> 18 master requirements, indicator->MR map,
                                                                       functional areas and platform modules
Outputs:
  supabase/seed/seed.sql   idempotent upserts for a Supabase project
  src/data/reference.json  the same records, bundled so the PWA works offline and in demo mode

Operational content that is not in the workbooks (hazards, protocol templates, AA checklist items,
offices) is defined below. It is DRAFT content for CDRRMO validation and is marked as such.

Usage: python3 scripts/generate_seed.py
"""
from __future__ import annotations
import json, re, uuid
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
IAT = ROOT / 'reference' / 'Ormoc_CDRRMO-EOC_Institutional_Assessment_Tool_v1.1.xlsx'
MRT = ROOT / 'reference' / 'Ormoc_CDRRMO_EOC_Master_Requirements_Traceability_and_140_Indicator_Data_Dictionary_v1.1.xlsx'
NS = uuid.UUID('6f1c1d1e-0a0b-4c4d-9e8f-0123456789ab')   # namespace for deterministic ids


def uid(*parts: str) -> str:
    return str(uuid.uuid5(NS, '/'.join(parts)))


# ----------------------------------------------------------------- workbooks --
def read_iat():
    wb = openpyxl.load_workbook(IAT)
    domains, indicators = [], []
    for sheet, part in (('A_CDRRMO', 'A'), ('B_EOC', 'B'), ('C_Interface', 'C')):
        ws = wb[sheet]
        for r in range(5, ws.max_row + 1):
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            if a and not b:
                code, rest = a.split('  ', 1)
                name, _, scope = rest.partition('  -  ')
                domains.append(dict(code=code, part=part, scope=scope.strip()))
            elif a:
                indicators.append(dict(
                    id=a, domain=b, statement=ws.cell(r, 3).value, reference=ws.cell(r, 4).value,
                    type='GP' if part == 'C' else ws.cell(r, 5).value, weight=int(ws.cell(r, 8).value),
                    means_of_verification=ws.cell(r, 6).value, methods=ws.cell(r, 7).value,
                    moo_chapter=ws.cell(r, 16).value,
                    lead_office=ws.cell(r, 5).value if part == 'C' else None))
    xw = wb['Crosswalk']
    crosswalk, framework = [], None
    for r in range(1, xw.max_row + 1):
        a = xw.cell(r, 1).value
        if not a:
            continue
        if str(a).startswith('A. UNDRR'):
            framework = 'UNDRR DRSC'; continue
        if str(a).startswith('B. RA 12287'):
            framework = 'RA 12287 IRR'; continue
        if framework and a not in ('DRSC ref', 'IRR provision'):
            crosswalk.append(dict(framework=framework, ref=str(a), subject=xw.cell(r, 2).value,
                                  treatment=xw.cell(r, 3).value, indicators=xw.cell(r, 4).value,
                                  note=xw.cell(r, 5).value))
    return domains, indicators, crosswalk


def read_mrt():
    wb = openpyxl.load_workbook(MRT)
    ws = wb['Master_Requirements']
    h = [c.value for c in ws[1]]
    mrs = [dict(zip(h, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    ws = wb['Domain_Traceability']
    h = [c.value for c in ws[1]]
    doms = {r[0]: dict(zip(h, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[0] != 'TOTAL'}
    ws = wb['Indicator_Data_Dictionary']
    h = [c.value for c in ws[1]]
    dd = {r[0]: dict(zip(h, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    return mrs, doms, dd


# --------------------------------------------------------- manual sections ----
MOO_TITLES = {
    'MOO Ch.1': 'Legal and Institutional Framework', 'MOO Ch.2': 'Office Organization',
    'MOO Ch.3': 'DRRM and Contingency Planning', 'MOO Ch.4': 'Operations, Warning and Response',
    'MOO Ch.5': 'DRRM Financing', 'MOO Ch.6': 'Capacity Development', 'MOO Ch.7': 'Knowledge Management, Reporting and M&E',
    'MOO Ch.8': 'Rehabilitation and Recovery', 'MOO Ch.9': 'Partnerships',
    'EOC Ch.1': 'Legal Basis', 'EOC Ch.2': 'Concept of Operations and Activation', 'EOC Ch.3': 'Facility, Site and Continuity',
    'EOC Ch.4': 'Communications and Sector Coordination', 'EOC Ch.5': 'Organization, Staffing and Public Information',
    'EOC Ch.6': 'Information Management and Reporting', 'EOC Ch.7': 'Coordination and Resource Management',
    'EOC Ch.8': 'Training and Exercises', 'EOC Ch.9': 'Sustainability',
}
ANNEX_JD = 'MOO Annex JDs'


def chapters_of(tag: str) -> list[str]:
    """'MOO Ch.2 / Ch.4' -> ['MOO Ch.2','MOO Ch.4']; 'EOC Manual (whole)' -> all EOC chapters."""
    if 'EOC Manual (whole)' in tag:
        return [k for k in MOO_TITLES if k.startswith('EOC')]
    out, prefix = [], None
    for part in [p.strip() for p in tag.split('/')]:
        m = re.match(r'^(MOO|EOC)\s+Ch\.(\d+)', part)
        if m:
            prefix = m.group(1); out.append(f'{prefix} Ch.{m.group(2)}'); continue
        m = re.match(r'^Ch\.(\d+)', part)
        if m and prefix:
            out.append(f'{prefix} Ch.{m.group(1)}'); continue
        if part.startswith('Annex JDs'):
            out.append(ANNEX_JD)
    return list(dict.fromkeys(out))


# ------------------------------------------------------------------ offices ---
OFFICES = [  # code, name, kind
    ('CDRRMO', 'City Disaster Risk Reduction and Management Office', 'cdrrmo'),
    ('CMO', "City Mayor's Office", 'city_office'),
    ('CSWDO', 'City Social Welfare and Development Office', 'city_office'),
    ('CHO', 'City Health Office', 'city_office'),
    ('CPDO', 'City Planning and Development Office', 'city_office'),
    ('CBO', 'City Budget Office', 'city_office'),
    ('CACCO', 'City Accounting Office', 'city_office'),
    ('BAC', 'Bids and Awards Committee', 'city_office'),
    ('CEO', 'City Engineering Office / Building Official', 'city_office'),
    ('GSO', 'General Services Office', 'city_office'),
    ('CENRO', 'City Environment and Natural Resources Office', 'city_office'),
    ('CAO', 'City Agriculture Office', 'city_office'),
    ('CVO', 'City Veterinary Office', 'city_office'),
    ('ZA', 'Zoning Administrator', 'city_office'),
    ('TRAFFIC', 'City traffic office', 'city_office'),
    ('MARKET', 'City market office', 'city_office'),
    ('PIO', 'City public information office', 'city_office'),
    ('DEPED-SDO', 'DepEd Schools Division of Ormoc City', 'national_agency'),
    ('DPWH', 'Department of Public Works and Highways (district office)', 'national_agency'),
    ('DENR-CENRO', 'DENR Community Environment and Natural Resources Office', 'national_agency'),
    ('DTI-LEYTE', 'DTI Leyte Provincial Office', 'national_agency'),
    ('OCD-VIII', 'Office of Civil Defense Region VIII / RDRRMC VIII', 'national_agency'),
    ('HOSPITALS', 'Hospitals in Ormoc City', 'other'),
    ('PORT', 'Port and terminal operators', 'private'),
    ('UTILITIES', 'Utility providers (power, water, telecommunications)', 'utility'),
    ('CHAMBER', 'Ormoc chamber of commerce', 'private'),
]
LEAD_MAP = {  # Part C lead-office text fragment -> office code
    'CDRRMO': 'CDRRMO', 'Chamber of Commerce': 'CHAMBER', 'utility providers': 'UTILITIES', 'CENRO': 'CENRO',
    'City Engineer': 'CEO', 'GSO': 'GSO', 'DENR CENRO': 'DENR-CENRO', 'CPDO': 'CPDO', 'City Budget Office': 'CBO',
    'Zoning Administrator': 'ZA', 'City Agriculture Office': 'CAO', 'City Veterinary Office': 'CVO',
    'City Engineer / Building Official': 'CEO', 'traffic office': 'TRAFFIC', 'DPWH': 'DPWH', 'port operator': 'PORT',
    'City Health Office': 'CHO', 'hospitals': 'HOSPITALS', 'DTI Leyte': 'DTI-LEYTE', 'city market office': 'MARKET',
    'DepEd Schools Division of Ormoc City': 'DEPED-SDO',
}


def lead_offices(text: str) -> list[str]:
    parts = []
    for chunk in re.split(r';', text):
        chunk = chunk.strip()
        for piece in re.split(r'\s+with\s+', chunk):
            piece = piece.strip()
            code = LEAD_MAP.get(piece)
            if code is None:
                raise SystemExit(f'Unmapped lead office: {piece!r} in {text!r}')
            parts.append(code)
    return list(dict.fromkeys(parts))


# ------------------------------------------------------------------ hazards ---
HAZARDS = [  # code, name, forecastable, typical lead h, declaration eligible, note
    ('tropical_cyclone', 'Tropical cyclone (severe winds)', True, 96, True, 'PAGASA bulletins; usually reaches the 3-day lead time'),
    ('storm_surge', 'Storm surge (Ormoc Bay)', True, 72, True, 'Forecast with tropical cyclones; coastal barangays'),
    ('flood', 'Riverine flood', True, 48, True, 'Eligible only when the forecast gives 3 days or more'),
    ('flash_flood', 'Flash flood (Anilao–Malbasag system)', True, 6, False, 'Lead time usually hours: city early action under LGU authority'),
    ('rain_landslide', 'Rain-induced landslide', True, 12, False, 'Upland barangays; city early action on rainfall thresholds'),
    ('drought', 'Drought (El Niño)', True, 720, True, 'Slow onset; NDRRMC to set timelines for drought'),
    ('extreme_heat', 'Extreme heat', True, 72, True, 'PAGASA heat index forecasts'),
    ('earthquake', 'Earthquake (Philippine Fault)', False, None, False, 'Not forecastable; outside RA 12287'),
    ('tsunami', 'Tsunami (local source)', False, None, False, 'Minutes of warning; outside RA 12287'),
    ('industrial', 'Industrial / geothermal incident', False, None, False, 'Not forecastable; outside RA 12287'),
]

# ------------------------------------------------- AA checklist (IRR Rule IV) -
AA_TEMPLATES = [  # sector, title, irr_ref, lead_office, timing, hazard
    ('Public information', 'Issue official notice of the declaration and the anticipatory actions being taken', 'IRR Rule IV Sec. 2(a)', 'CDRRMO / public information office', 'on_declaration', None),
    ('Public information', 'Disseminate impact-based advisories on all channels, including GIDA barangays, in Cebuano and accessible formats', 'IRR Rule IV Sec. 2(b)-(c)', 'CDRRMO / public information office', 'before_impact_beyond_3d', None),
    ('Prepositioning', 'Preposition food and non-food relief items for the forecast affected population', 'IRR Rule IV Sec. 3(a)(i); Sec. 6(a)', 'CSWDO', 'before_impact_beyond_3d', None),
    ('Prepositioning', 'Preposition essential medicines, medical supplies and hygiene kits', 'IRR Rule IV Sec. 3(a)(ii)', 'CHO', 'before_impact_beyond_3d', None),
    ('Prepositioning', 'Stock animal feed, forage, emergency water and mobile veterinary kits', 'IRR Rule IV Sec. 3(a)(iii)', 'City Veterinary Office', 'within_3d', None),
    ('Prepositioning', 'Secure water and power supply for evacuation centres', 'IRR Rule IV Sec. 3(a)(iv)', 'GSO / utility providers', 'within_3d', None),
    ('Operational assets', 'Pre-deploy heavy equipment and road-clearing convoys', 'IRR Rule IV Sec. 3(b)(i)', 'City Engineering Office', 'within_3d', None),
    ('Operational assets', 'Activate transport for pre-emptive evacuation and supply distribution', 'IRR Rule IV Sec. 3(b)(ii)', 'GSO / CDRRMO', 'within_3d', None),
    ('Operational assets', 'Preposition response teams and their equipment', 'IRR Rule IV Sec. 3(b)(iii)', 'CDRRMO', 'within_3d', None),
    ('Evacuation', 'Conduct pre-emptive or forced evacuation of at-risk families', 'RA 12287 Sec. 9(c); IRR Rule IV Sec. 4(a)(i)', 'CDRRMO / BDRRMCs', 'within_3d', None),
    ('Evacuation', 'Assist evacuation of livestock, pets, crops and livelihood assets', 'IRR Rule IV Sec. 4(a)(ii)', 'City Agriculture Office / City Veterinary Office', 'within_3d', None),
    ('Evacuation', 'Open safe and inclusive evacuation centres and apply CCCM and IDP protection protocols', 'IRR Rule IV Sec. 4(b)', 'CSWDO', 'within_3d', None),
    ('Lifelines', 'Rapid safety inspection of evacuation centres and other critical facilities', 'IRR Rule IV Sec. 4(c)(i)', 'Building Official', 'within_3d', None),
    ('Lifelines', 'Coordinate load-shedding, water rationing and alternate logistics with utilities', 'IRR Rule IV Sec. 4(c)(ii)-(iii)', 'Utility providers', 'within_3d', None),
    ('Lifelines', 'Request pre-deployment of cell-on-wheels or satellite communication units', 'IRR Rule IV Sec. 4(c)(iv)', 'CDRRMO / telecommunications providers', 'within_3d', None),
    ('Volunteers', 'Mobilize accredited community disaster volunteers, CSOs and private partners under the deployment plan', 'RA 12287 Sec. 9(d); IRR Rule IV Sec. 5', 'CDRRMO', 'within_3d', None),
    ('Social amelioration', 'Identify and prioritize families using CBMS or the prevailing registry (no personal data copied into this platform)', 'IRR Rule IV Sec. 7(a)', 'CSWDO', 'before_impact_beyond_3d', None),
    ('Social amelioration', 'Release anticipatory cash or vouchers under the City guidelines', 'IRR Rule IV Sec. 7(b)-(c)', 'CSWDO / City Treasurer', 'within_3d', None),
    ('Agriculture and fisheries', 'Activate early-action protocols for farmers and fisherfolk (RSBSA, FishR, livestock registries)', 'IRR Rule IV Sec. 8', 'City Agriculture Office', 'before_impact_beyond_3d', None),
    ('Health', 'Activate the Health Cluster, service-continuity and surge plans; put health teams on standby', 'IRR Rule IV Sec. 9(a)', 'CHO / hospitals', 'within_3d', None),
    ('Education', 'Activate school anticipatory action plans and learning-continuity modalities', 'IRR Rule IV Sec. 10(a)', 'DepEd Schools Division', 'within_3d', None),
    ('Transportation', 'Pre-inspect bridges, clear roads, set alternate routes and traffic teams; coordinate with port operators', 'IRR Rule IV Sec. 10(b)', 'City Engineering Office / traffic office / DPWH', 'within_3d', None),
    ('Economy', 'Monitor prices of basic goods and act on profiteering before impact', 'IRR Rule IV Sec. 10(c)', 'DTI Leyte / city market office', 'within_3d', None),
    ('Environment', 'Declog waterways and drainage; stabilize slopes and coasts in delineated hazard zones', 'IRR Rule IV Sec. 10(d)(i)-(ii)', 'CENRO / City Engineering Office', 'before_impact_beyond_3d', None),
    ('Environment', 'Prepare safe disposal of debris and livestock carcasses', 'IRR Rule IV Sec. 10(d)(iii)', 'CENRO / GSO', 'within_3d', None),
    ('Evacuation', 'Evacuate coastal puroks exposed to the forecast storm surge height', 'IRR Rule IV Sec. 4(a)', 'CDRRMO / coastal BDRRMCs', 'within_3d', 'storm_surge'),
    ('Agriculture and fisheries', 'Advise early harvest and safe storage of crops; secure fishing boats', 'IRR Rule IV Sec. 8(b)', 'City Agriculture Office', 'before_impact_beyond_3d', 'tropical_cyclone'),
    ('Water supply', 'Implement water rationing and protect drinking-water sources', 'IRR Rule IV Sec. 4(c)(ii)', 'Water service provider / CHO', 'before_impact_beyond_3d', 'drought'),
    ('Reporting', 'Submit the report to OCD within 45 days of lifting: declaration status, AA fund use, estimated costs minimized', 'RA 12287 Sec. 8; IRR Rule VI Sec. 1', 'CDRRMO', 'after_lifting', None),
]

# --------------------------------------------------------- CDRRMO protocols ---
# Draft templates for CDRRMO validation. Thresholds that the City has not yet
# set are written as instructions to apply the City's own matrix, not as numbers.
PROTOCOLS = [
    dict(code='PR-DUTY-HANDOVER', title='Duty shift handover', category='daily_operations', owner_role='duty_officer',
         doc='MOO Ch.4', purpose='Keep 24/7 operations continuous between shifts.',
         trigger='Every change of duty shift at the Operations Center.',
         steps=[('Read the logbook entries since the last handover', False, False, 0),
                ('Review open incidents, resource requests and pending barangay reports', True, False, 5),
                ('Check the latest PAGASA and PHIVOLCS bulletins and record any change', True, True, 10),
                ('Run the radio net check with the barangays scheduled for this shift', True, True, 20),
                ('Check generator fuel, UPS and backup internet status', False, True, 25),
                ('Record handover notes and sign off with the incoming duty officer', True, True, 30)]),
    dict(code='PR-WARNING-DISSEM', title='Warning dissemination', category='warning', owner_role='duty_officer',
         doc='EOC Ch.2', purpose='Move a hazard bulletin to every at-risk purok, in a form people understand.',
         trigger='On receipt of a PAGASA, PHIVOLCS or MGB bulletin or advisory that affects Ormoc.',
         steps=[('Log the bulletin: issuer, number, time issued, hazard and areas', True, True, 0),
                ('Check it against the City warning-to-action matrix and note the matching alert and actions', True, True, 10),
                ('Notify the EOC Manager and CDRRMO Head; record the instruction received', True, True, 15),
                ('Issue the advisory in Cebuano, with Filipino/English and accessible formats as needed', True, False, 30),
                ('Send through every channel: emergency cell broadcast/SMS, radio, social media, barangay officials', True, True, 40),
                ('Confirm receipt from riverine, coastal and upland barangays first, then the rest', True, True, 60),
                ('Log barangays that did not confirm and assign follow-up', False, True, 75)]),
    dict(code='PR-EOC-ACTIVATE', title='EOC activation (Blue or Red)', category='eoc', owner_role='eoc_manager',
         doc='EOC Ch.2', purpose='Stand up the EOC quickly with the right people and records.',
         trigger='When the EOC Manager raises the level to Blue or Red.',
         steps=[('Record the level change and reason in the platform (logs it automatically). If no device can connect, write it in the paper logbook with the time and enter it once connected', True, False, 0),
                ('Notify CDRRMC members, cluster leads and liaison desks of the activation', True, True, 15),
                ('Fill EOC positions from the roster, with alternates', True, True, 30),
                ('Run the communications check on all redundant channels', True, True, 30),
                ('Set the operational period and the SitRep schedule', True, True, 45),
                ('Brief the Mayor on the situation and decisions needed', True, True, 60)]),
    dict(code='PR-PDRA', title='Pre-disaster risk assessment (PDRA)', category='anticipatory_action', owner_role='eoc_manager',
         doc='EOC Ch.2', purpose='Decide, from the forecast and local exposure, whether to monitor, act early or request a declaration.',
         trigger='When a forecast hazard may affect Ormoc.',
         steps=[('Convene the PDRA core group and record attendance by office', True, False, 0),
                ('Review the forecast: hazard, expected impact time, areas, impact level', True, True, 15),
                ('Identify exposed barangays, vulnerable groups, critical facilities and possible cascading hazards', True, True, 30),
                ('Classify the projected impact and compute the lead time left', True, False, 45),
                ('Decide: monitor, city early action, or request a State of Imminent Disaster; record the basis', True, True, 60),
                ('Record the PDRA in the platform and brief the CDRRMO Head', True, False, 75)]),
    dict(code='PR-SOID-REQUEST', title='State of Imminent Disaster: request and declaration', category='anticipatory_action', owner_role='eoc_manager',
         doc='MOO Ch.1', purpose='Follow RA 12287 and its IRR from the PDRA to the Mayor\'s Executive Order.',
         trigger='When the PDRA shows severe impact with at least 3 days of lead time.',
         steps=[('Prepare the CDRRMC request to RDRRMC VIII with the PDRA results', True, False, 0),
                ('Send the request and record the time sent (RDRRMC replies within 24 hours)', True, True, 15),
                ('Record the RDRRMC recommendation and its reference', True, True, None),
                ('Prepare the Mayor\'s Executive Order for signature', True, False, None),
                ('Record the EO number and time of declaration', True, True, None),
                ('Issue official public notice of the declaration and the actions required', True, False, None),
                ('Load the anticipatory action checklist and assign lead offices', True, False, None)]),
    dict(code='PR-FLASHFLOOD-EARLY', title='Flash flood and landslide early action', category='anticipatory_action', owner_role='duty_officer',
         doc='EOC Ch.2', purpose='Act on short-lead hazards that cannot wait for a declaration.',
         trigger='When rainfall or river readings reach the City\'s early-action thresholds for the Anilao–Malbasag system or upland barangays.',
         steps=[('Record the reading, gauge or source, and the threshold reached', True, True, 0),
                ('Alert riverine and upland BDRRMCs by radio and SMS', True, True, 5),
                ('Recommend pre-emptive evacuation of the identified puroks to the EOC Manager', True, True, 10),
                ('Open the designated evacuation centres and alert CSWDO', True, False, 20),
                ('Position rescue teams at the agreed staging points', True, True, 30),
                ('Confirm evacuation counts from each barangay and log them', True, True, 60)]),
    dict(code='PR-SITREP-CYCLE', title='SitRep cycle', category='reporting', owner_role='eoc_staff',
         doc='EOC Ch.6', purpose='Produce a verified SitRep on schedule every operational period.',
         trigger='At each scheduled SitRep time during an activation.',
         steps=[('Collect barangay reports and cluster updates since the last SitRep', True, False, 0),
                ('Reconcile evacuation figures with evacuation-centre counts', True, True, 20),
                ('Draft the SitRep in the standard format', True, False, 40),
                ('EOC Manager reviews and approves', True, False, 55),
                ('Submit to OCD VIII and the Province; record the time sent', True, True, 60)]),
    dict(code='PR-CDRRMC-MEETING', title='CDRRMC meeting (secretariat)', category='council', owner_role='duty_officer',
         doc='MOO Ch.1', purpose='Run the Council meeting and keep its records complete.',
         trigger='Quarterly, and whenever the Chair calls a special meeting.',
         steps=[('Send the notice and agenda to all members', True, True, None),
                ('Record attendance and confirm quorum', True, True, None),
                ('Record decisions and draft resolutions', True, True, None),
                ('Circulate the minutes for approval', False, False, None),
                ('File approved minutes and resolutions; post as required', True, False, None)]),
    dict(code='PR-LIFT-REPORT', title='Lifting and 45-day report', category='anticipatory_action', owner_role='eoc_manager',
         doc='MOO Ch.7', purpose='Lift a State of Imminent Disaster correctly and report to OCD on time.',
         trigger='When the hazard occurs or the PDRA shows the projected impacts will not be met.',
         steps=[('Record the lifting ground and its basis (advisory, incident or situation report, or latest PDRA)', True, True, 0),
                ('Shift to response and early recovery if the hazard occurred', True, False, None),
                ('If not utilized, return LDRRMF anticipatory funds to the Special Trust Fund and hand stocks to CSWDO', True, True, None),
                ('Deliver perishable goods to the intended beneficiaries (no-regrets principle)', False, True, None),
                ('Compile declaration status, AA fund use and estimated costs minimized', True, False, None),
                ('Submit the report to OCD within 45 days of lifting and record the reference', True, True, None)]),
]


# ------------------------------------------------------------------- build ----
def main():
    domains, indicators, crosswalk = read_iat()
    mrs, doms_mrt, dd = read_mrt()
    assert len(indicators) == 140, len(indicators)

    out = dict(parts=[dict(code='A', name='Part A – CDRRMO institutional capacity', in_imi=True, sort=1),
                      dict(code='B', name='Part B – Emergency Operations Center', in_imi=True, sort=2),
                      dict(code='C', name='Part C – Interface indicators', in_imi=False, sort=3)])
    out['domains'] = []
    for i, d in enumerate(domains, 1):
        m = doms_mrt[d['code']]
        out['domains'].append(dict(code=d['code'], part=d['part'], name=m['Domain'], scope=d['scope'],
                                   functional_area=m['Primary Functional Area'], platform_module=m['Platform Module'], sort=i))
    out['indicators'] = []
    for i, ind in enumerate(indicators, 1):
        tags = [t.strip() for t in str(dd[ind['id']]['Thematic Tags*']).split(';') if t.strip() and t.strip() != '—']
        out['indicators'].append(dict(ind, tags=tags, version_status=dd[ind['id']]['Version Status (v1.1)'], sort=i))
    out['master_requirements'] = [dict(
        id=m['Req ID'], title=m['Requirement / Source Objective'], statement=m['Requirement Statement'],
        source_basis=m['Source Basis'], iat_coverage=m['IAT Coverage'], evidence=m['Evidence / Verification'],
        sa_output=m['Situational Analysis Output'], moo_output=m['Draft MOO Output'],
        platform_req=m['Platform Requirement / UX Workflow'], data_entities=m['Core Data Entities'],
        escalation_rule=m['Priority / Escalation Rule'], priority=m['Priority'], sort=i) for i, m in enumerate(mrs, 1)]
    links = []
    for ind in indicators:
        raw = str(dd[ind['id']]['Master Requirement(s)*'])
        ids = []
        for tok in [t.strip() for t in raw.split(';')]:
            if tok == 'MR-12–MR-15':
                ids += ['MR-12', 'MR-13', 'MR-14', 'MR-15']
            elif tok:
                ids.append(tok)
        for mr in dict.fromkeys(ids):
            links.append(dict(indicator_id=ind['id'], mr_id=mr))
    out['indicator_requirements'] = links
    out['crosswalk'] = [dict(c, sort=i) for i, c in enumerate(crosswalk, 1)]

    out['offices'] = [dict(id=uid('office', c), code=c, name=n, kind=k) for c, n, k in OFFICES]
    office_id = {o['code']: o['id'] for o in out['offices']}
    out['indicator_lead_offices'] = [dict(indicator_id=i['id'], office_id=office_id[c])
                                     for i in indicators if i['lead_office'] for c in lead_offices(i['lead_office'])]

    out['hazards'] = [dict(code=c, name=n, forecastable=f, typical_lead_hours=l, declaration_eligible=e, note=no, sort=i)
                      for i, (c, n, f, l, e, no) in enumerate(HAZARDS, 1)]

    # Manual sections: one skeleton per chapter listing the indicators it must answer.
    docs, doc_links, versions = [], [], []
    by_ch: dict[str, list[dict]] = {}
    for ind in indicators:
        for ch in chapters_of(ind['moo_chapter']):
            by_ch.setdefault(ch, []).append(ind)
    codes = list(MOO_TITLES) + [ANNEX_JD]
    missing = [c for c in by_ch if c not in codes]
    assert not missing, missing
    for i, code in enumerate(codes, 1):
        manual = code.split()[0]
        title = MOO_TITLES.get(code, 'Job descriptions')
        did = uid('doc', code)
        inds = by_ch.get(code, [])
        body = [f'# {code} – {title}', '',
                '> Skeleton generated from IAT v1.1. Replace each requirement with the City\'s procedure, standard or form.', '',
                '## Requirements this section must meet', '']
        for ind in inds:
            body.append(f'- **{ind["id"]}** ({ind["type"]}, weight {ind["weight"]}) {ind["statement"]}  ')
            body.append(f'  *Basis:* {ind["reference"]}')
        vid = uid('docver', code, 'v0.1')
        docs.append(dict(id=did, manual=manual, kind='annex' if code == ANNEX_JD else 'chapter', code=code, title=title,
                         summary=f'{len(inds)} IAT indicator(s) route their gaps here.', status='draft', sort=i,
                         current_version_id=None))
        versions.append(dict(id=vid, document_id=did, label='v0.1 skeleton', body_md='\n'.join(body),
                             change_note='Generated from IAT v1.1', status='draft'))
        for ind in inds:
            doc_links.append(dict(document_id=did, indicator_id=ind['id']))
    out['moo_documents'], out['moo_versions'], out['moo_document_indicators'] = docs, versions, doc_links

    tpls, steps = [], []
    for p in PROTOCOLS:
        tid = uid('protocol', p['code'])
        tpls.append(dict(id=tid, code=p['code'], title=p['title'], purpose=p['purpose'], trigger_text=p['trigger'],
                         document_id=uid('doc', p['doc']), owner_role=p['owner_role'], category=p['category'],
                         status='draft', template_rev=1))
        for seq, (instr, crit, note, due) in enumerate(p['steps'], 1):
            steps.append(dict(id=uid('pstep', p['code'], str(seq)), template_id=tid, seq=seq, instruction=instr,
                              is_critical=crit, requires_note=note, due_offset_minutes=due))
    out['protocol_templates'], out['protocol_steps'] = tpls, steps

    out['aa_action_templates'] = [dict(id=uid('aa', str(i)), hazard_code=h, sector=s, title=t, irr_ref=r, lead_office=lo,
                                       timing=tm, seq=i)
                                  for i, (s, t, r, lo, tm, h) in enumerate(AA_TEMPLATES, 1)]

    (ROOT / 'src' / 'data').mkdir(parents=True, exist_ok=True)
    (ROOT / 'src' / 'data' / 'reference.json').write_text(json.dumps(out, ensure_ascii=False, indent=1))
    write_sql(out)
    print({k: len(v) for k, v in out.items()})


def lit(v):
    if v is None:
        return 'null'
    if isinstance(v, bool):
        return 'true' if v else 'false'
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, list):
        return "array[" + ','.join(lit(x) for x in v) + "]::text[]" if v else "'{}'::text[]"
    return "'" + str(v).replace("'", "''") + "'"


def upsert(table, rows, key, cols=None, cast=None):
    if not rows:
        return ''
    cols = cols or list(rows[0].keys())
    cast = cast or {}
    vals = ',\n'.join('(' + ','.join(lit(r.get(c)) + cast.get(c, '') for c in cols) + ')' for r in rows)
    keys = key if isinstance(key, list) else [key]
    upd = [c for c in cols if c not in keys]
    action = ('do update set ' + ', '.join(f'{c} = excluded.{c}' for c in upd)) if upd else 'do nothing'
    return f'insert into public.{table} ({",".join(cols)}) values\n{vals}\non conflict ({",".join(keys)}) {action};\n\n'


def write_sql(o):
    sql = ['-- GENERATED by scripts/generate_seed.py from reference/*.xlsx — do not edit by hand.\n',
           '-- Reference data only. No personal data. Protocols and AA items are DRAFT for CDRRMO validation.\n',
           'begin;\n\n']
    sql.append(upsert('iat_parts', o['parts'], 'code'))
    sql.append(upsert('iat_domains', o['domains'], 'code'))
    sql.append(upsert('iat_indicators', o['indicators'], 'id'))
    sql.append(upsert('master_requirements', o['master_requirements'], 'id'))
    sql.append('delete from public.indicator_requirements;\n')
    sql.append(upsert('indicator_requirements', o['indicator_requirements'], ['indicator_id', 'mr_id']))
    sql.append('delete from public.framework_crosswalk;\n')
    cols = ['framework', 'ref', 'subject', 'treatment', 'indicators', 'note', 'sort']
    sql.append('insert into public.framework_crosswalk (' + ','.join(cols) + ') values\n' +
               ',\n'.join('(' + ','.join(lit(c.get(k)) for k in cols) + ')' for c in o['crosswalk']) + ';\n\n')
    sql.append(upsert('offices', o['offices'], 'id', cast={'id': '::uuid'}))
    sql.append(upsert('indicator_lead_offices', o['indicator_lead_offices'], ['indicator_id', 'office_id'], cast={'office_id': '::uuid'}))
    sql.append(upsert('hazards', o['hazards'], 'code'))
    sql.append(upsert('moo_documents', o['moo_documents'], 'id',
                      cols=['id', 'manual', 'kind', 'code', 'title', 'summary', 'status', 'sort'], cast={'id': '::uuid'}))
    sql.append(upsert('moo_versions', o['moo_versions'], 'id', cast={'id': '::uuid', 'document_id': '::uuid'}))
    sql.append(upsert('moo_document_indicators', o['moo_document_indicators'], ['document_id', 'indicator_id'], cast={'document_id': '::uuid'}))
    sql.append(upsert('protocol_templates', o['protocol_templates'], 'id',
                      cast={'id': '::uuid', 'document_id': '::uuid', 'owner_role': '::public.app_role'}))
    sql.append(upsert('protocol_steps', o['protocol_steps'], 'id', cast={'id': '::uuid', 'template_id': '::uuid'}))
    sql.append(upsert('aa_action_templates', o['aa_action_templates'], 'id', cast={'id': '::uuid'}))
    sql.append('commit;\n')
    (ROOT / 'supabase' / 'seed').mkdir(parents=True, exist_ok=True)
    (ROOT / 'supabase' / 'seed' / 'seed.sql').write_text(''.join(sql))


if __name__ == '__main__':
    main()
