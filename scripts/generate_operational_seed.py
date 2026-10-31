#!/usr/bin/env python3
"""Generate one-time operational seed files from the City's source documents.

Unlike scripts/generate_seed.py (reference data, safe to re-run at every release), these files load
working records that the City then owns and edits in the app. Run each file once, after seed.sql,
in the Supabase SQL editor (or psql). Each file is safe to run again: it never overwrites a record
that already exists, and never changes a score that someone has set in the app.

Inputs  (reference/):
  Ormoc_CDRRMO-EOC_Institutional_Assessment_Tool_v*.xlsx   October 2026 baseline: scores, evidence,
                                                           gaps, recommended actions, Action_Plan, Profile
Outputs (supabase/seed/operational/):
  01-03_iat_baseline_*.sql   assessment cycle "2026 Baseline" with 140 provisional responses, their
                             evidence notes, scores and the improvement action plan (three parts,
                             each small enough for the Supabase SQL editor; run in order)
  04_vault_register.sql      Document Vault entries for the source documents used in the engagement

Evidence rule (IAT Methodology 5.2; app.enforce_evidence_gate): interviews and FGDs alone cannot
support a score above 2; a 3 needs an approved document and proof of use; a 4 also needs a test,
exercise or after-action review. The workbook records field self-scores ("F:") and desk review
("DR:"). This import keeps that rule:
  F:  -> evidence kind 'fgd', strength 'verbal' (caps at 2)
  DR: -> evidence kind 'document'; strength 'approved_in_use' when the assessor's score is 3 or 4,
         otherwise 'partial'. The note says the strength must be confirmed against the document.
  No source shows a test or after-action review, so no score is imported above 3.
Where the workbook score is higher than the evidence supports, the response is imported at the
supported level and Gap / Remarks starts with the workbook score and the reason. Restore it in the
app once the evidence is attached; the evidence gate will then accept it.

Usage: python3 scripts/generate_operational_seed.py
"""
from __future__ import annotations
import json, re, uuid
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'supabase' / 'seed' / 'operational'
NS = uuid.UUID('6f1c1d1e-0a0b-4c4d-9e8f-0123456789ab')   # same namespace as generate_seed.py
CYCLE_KEY = '2026-baseline'
PART = '--@@PART@@'


def uid(*parts: str) -> str:
    return str(uuid.uuid5(NS, '/'.join(parts)))


def one(pattern: str) -> Path:
    found = sorted((ROOT / 'reference').glob(pattern))
    if len(found) != 1:
        raise SystemExit(f'Expected exactly one reference/{pattern}, found {[f.name for f in found]}')
    return found[0]


def lit(v) -> str:
    if v is None:
        return 'null'
    if isinstance(v, bool):
        return 'true' if v else 'false'
    if isinstance(v, (int, float)):
        return str(int(v)) if float(v).is_integer() else str(v)
    if isinstance(v, (list, tuple)):
        return "array[" + ','.join(lit(x) for x in v) + "]::text[]"
    t = str(v)
    if '\n' in t or '\r' in t:   # keep each row on one line (SQL editor line limit)
        t = t.replace('\\', '\\\\').replace("'", "''").replace('\r', '').replace('\n', '\\n')
        return "E'" + t + "'"
    return "'" + t.replace("'", "''") + "'"


# Positions only (RA 10173; workbook Profile): drop personal names that appear in the evidence text.
REDACT = [(re.compile(r'\s*\(Ciriaco Tolibao\)'), '')]


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    for pat, rep in REDACT:
        s = pat.sub(rep, s)
    return None if s in ('', '—', '-') else s


IND_ID = re.compile(r'^[ABC]\d+\.\d+$')
SEG = re.compile(r'(?:(?<=^)|(?<=[\s.;]))(F|DR)\s*(\([^)]*\))?:\s*')


def segments(text: str):
    """Split 'F: ... DR: ...' into [('F', '...'), ('DR', '...')]."""
    marks = list(SEG.finditer(text))
    out = []
    for k, m in enumerate(marks):
        end = marks[k + 1].start() if k + 1 < len(marks) else len(text)
        body = (((m.group(2) or '') + ' ') if m.group(2) else '') + text[m.end():end].strip()
        if body.strip():
            out.append((m.group(1), body.strip()))
    return out


def read_workbook():
    wb = openpyxl.load_workbook(one('Ormoc_CDRRMO-EOC_Institutional_Assessment_Tool_v*.xlsx'))
    rows, where = {}, {}
    for sheet in ('A_CDRRMO', 'B_EOC', 'C_Interface'):
        ws = wb[sheet]
        for r in range(5, ws.max_row + 1):
            ind = ws.cell(r, 1).value
            if not (ind and IND_ID.match(str(ind))):
                continue
            score = ws.cell(r, 9).value
            rows[ind] = dict(
                id=ind, score=score, evidence=clean(ws.cell(r, 10).value) or '',
                gap=clean(ws.cell(r, 11).value), action=clean(ws.cell(r, 12).value))
            where[(sheet, r)] = ind
    plan = {}
    ws = wb['Action_Plan']
    ref = re.compile(r"^=(A_CDRRMO|B_EOC|C_Interface)!A(\d+)$")
    for r in range(4, ws.max_row + 1):
        m = ref.match(str(ws.cell(r, 1).value or ''))
        if not m:
            continue
        ind = where[(m.group(1), int(m.group(2)))]
        plan[ind] = dict(responsible=clean(ws.cell(r, 8).value), target=clean(ws.cell(r, 9).value),
                         budget=clean(ws.cell(r, 10).value), status=clean(ws.cell(r, 12).value))
    profile = {}
    for a, b in wb['Profile'].iter_rows(min_col=1, max_col=2, values_only=True):
        if a and b:
            profile[str(a).strip()] = str(b).strip()
    return rows, plan, profile


STATUS = {'Not started': 'not_started', 'Ongoing': 'ongoing', 'Complete': 'done', 'Completed': 'done',
          'Done': 'done', 'Dropped': 'dropped'}


def baseline_sql(rows, plan, profile, known_ids):
    missing = sorted(set(known_ids) - set(rows))
    extra = sorted(set(rows) - set(known_ids))
    assert not missing and not extra, (missing, extra)
    cyc = uid('cycle', CYCLE_KEY)
    notes = '\n'.join(f'{k}: {profile[k]}' for k in (
        'Assessment period (from - to)', 'FGD groups conducted', 'Sites inspected', 'Documents reviewed (count)',
        'Validation workshop date', 'Number of barangays') if k in profile)
    notes = ('Imported from the IAT v1.1 workbook (file v1-3, desk review completed 30 Oct 2026). '
             'Scores are provisional until the validation workshop.\n' + notes)
    out = ['-- GENERATED by scripts/generate_operational_seed.py from reference/Ormoc_CDRRMO-EOC_Institutional_Assessment_Tool_v*.xlsx',
           '-- One-time load of the October 2026 IAT baseline. Run after seed.sql. Safe to run again: existing',
           '-- responses, evidence and actions are left alone, and a score already set in the app is never changed.',
           '-- No personal data: evidence sources are offices and sections, never informants.',
           'begin;', '',
           'insert into public.assessment_cycles (id, name, cycle_year, iat_version, status, period_from, period_to, notes)',
           f"values ({lit(cyc)}::uuid, '2026 Baseline (IAT v1.1)', 2026, '1.1', 'fieldwork', '2026-10-01', '2026-10-30', {lit(notes)})",
           'on conflict (id) do nothing;', '']
    resp, ev, upd, acts = [], [], [], []
    held = []
    for ind in known_ids:
        r = rows[ind]
        rid = uid('response', CYCLE_KEY, ind)
        wb_score = None if r['score'] in (None, '') else r['score']
        is_na = isinstance(wb_score, str) and wb_score.upper() == 'NA'
        wb_score = None if is_na else (int(round(float(wb_score))) if wb_score is not None else None)
        segs = segments(r['evidence'])
        has_dr = any(k == 'DR' for k, _ in segs)
        cap = 0
        for n, (kind, text) in enumerate(segs, 1):
            if kind == 'F':
                strength, k, title, source = 'verbal', 'fgd', 'CDRRMO field self-assessment (FGD), October 2026', 'CDRRMO section FGDs'
                note = text
            else:
                strong = wb_score is not None and wb_score >= 3
                strength = 'approved_in_use' if strong else 'partial'
                k, title, source = 'document', 'Desk review, October 2026', 'Consultant desk review'
                note = text + (' [Imported at "approved and in use" from the assessor\'s desk score. Attach the document '
                               'and proof of use, or lower the strength, at validation.]' if strong else '')
            cap = max(cap, {'verbal': 2, 'partial': 2, 'approved_in_use': 3}[strength])
            ev.append(f"({lit(uid('evidence', CYCLE_KEY, ind, str(n)))}::uuid, {lit(rid)}::uuid, {lit(k)}, "
                      f"{lit(strength)}::public.evidence_strength, {lit(title)}, {lit(source)}, {lit(note)})")
        score = wb_score
        gap = r['gap']
        if wb_score is not None and wb_score > 0 and wb_score > cap:
            score = cap if cap else None
            why = ('it rests on FGD self-scoring only (interviews and FGDs cannot support more than 2)' if not has_dr
                   else 'no test, exercise or after-action review is on record (a 4 needs one)')
            prefix = (f'[Workbook score {wb_score}; imported as {score} because {why}. '
                      f'Restore it once the evidence is attached.]')
            gap = prefix + (' ' + gap if gap else '')
            held.append((ind, wb_score, score, 'F-only' if not has_dr else 'no test'))
        resp.append(f"({lit(rid)}::uuid, {lit(cyc)}::uuid, {lit(ind)}, {lit(is_na)}, "
                    f"{lit('Marked NA in the workbook' if is_na else None)}, {lit(gap)}, {lit(r['action'])}, 'provisional')")
        if score is not None:
            upd.append(f"({lit(rid)}::uuid, {score})")
        p = plan.get(ind)
        if r['action'] and p:
            text = r['action'] + (f"\n\nTarget (from the Action Plan): {p['target']}" if p['target'] else '')
            acts.append(f"({lit(uid('action', CYCLE_KEY, ind))}::uuid, {lit(cyc)}::uuid, {lit(ind)}, {lit(text)}, "
                        f"{lit(p['responsible'])}, {lit(p['budget'])}, {lit(STATUS.get(p['status'] or '', 'not_started'))})")

    out += ['-- Responses first, without scores: the evidence gate needs the evidence in place before a score.',
            'insert into public.responses (id, cycle_id, indicator_id, is_na, na_justification, gap_remarks, recommended_action, status) values',
            ',\n'.join(resp), 'on conflict (cycle_id, indicator_id) do nothing;', '',
            'commit;', PART, 'begin;',
            'insert into public.evidence (id, response_id, kind, strength, title, source, note) values',
            ',\n'.join(ev), 'on conflict (id) do nothing;', '',
            'commit;', PART, 'begin;',
            '-- Scores, only where no one has scored the indicator in the app yet.',
            'update public.responses r set score = v.score::smallint',
            'from (values', ',\n'.join(upd), ') as v(id, score)',
            'where r.id = v.id and r.score is null and not r.is_na;', '',
            'insert into public.actions (id, cycle_id, indicator_id, action, responsible_person, budget_source, status) values',
            ',\n'.join(acts), 'on conflict (id) do nothing;', '',
            'commit;', '',
            '-- Check: 140 responses, scored count, held scores, actions; then the cycle results.',
            f"select count(*) as responses, count(score) as scored, count(*) filter (where gap_remarks like '[Workbook score%') as held_below_workbook",
            f"  from public.responses where cycle_id = {lit(cyc)}::uuid;",
            f"select count(*) as actions from public.actions where cycle_id = {lit(cyc)}::uuid;",
            f"select grp, label, scored, maturity_index, maturity_band, mandatory_compliance, high_gaps",
            f"  from public.v_cycle_results where cycle_id = {lit(cyc)}::uuid order by grp;", '']
    return '\n'.join(out), held, len(acts), len(ev)


# Document Vault: the source documents of the engagement. Files are attached in the app
# (Vault -> document -> Attach file); originals are in the repository's reference/ folder.
VAULT = [  # key, title, category, ref_key, doc_number, issuer, date, status, indicators, url, location, summary
    ('ra12287-irr', 'Implementing Rules and Regulations of RA 12287 (State of Imminent Disaster)', 'irr', 'RA 12287 IRR',
     None, 'NDRRMC / Office of Civil Defense', '2026-09-14', 'current', ['A1.8', 'A2.9', 'A6.7', 'A10.8'], None,
     'reference/IRR_12287_2026.pdf', 'Certified IRR. Basis for the PDRA, declaration, anticipatory action, fund rules and the 45-day report.'),
    ('ndrrmc-memo-131', 'NDRRMC Memo No. 131, s. 2018: Guidelines on the establishment, operationalization and management of EOCs',
     'issuance', 'NDRRMC EOC Guidelines', 'Memo No. 131, s. 2018', 'NDRRMC', None, 'current',
     ['B1.1', 'B1.2', 'B1.3', 'B1.4', 'B4.1', 'B7.1'], None,
     'reference/NDRRMC Memo 131 s 2018 Guidelines on the Establishment Operationalization and Management of EOC.pdf',
     'National EOC guidelines: establishment, activation, organization and management.'),
    ('dilg8-reporting', 'DILG Region VIII memo on the Disaster Response Reporting Template and SOP adoption (RDRRMC8 Res. 12-006)',
     'issuance', None, 'RDRRMC VIII Res. No. 12-006', 'DILG Region VIII / RDRRMC VIII', '2025-02-26', 'current',
     ['A10.3', 'B1.2', 'B5.1', 'B5.2'], None,
     'reference/Disaster Response Reporting Template and SOP Adoption_DILG8.pdf',
     'Regional reporting template, schedule and EOC alert-level SOP.'),
    ('drrm-10yr-plan', 'Ormoc City DRRM 10-Year Plan 2023-2032', 'plan', None, None, 'City Government of Ormoc',
     None, 'current', ['A3.1', 'A3.5', 'A10.4'], None, 'reference/Ormoc City DRRM 10-Year Plan (2023-2032).pdf',
     'Local DRRM plan, including the 10-year LDRRMF investment plan and the results-based M&E design.'),
    ('pscp-guidebook', 'Public Service Continuity Plan Guidebook', 'issuance', 'PSCP Guidebook', None, 'NDRRMC / OCD',
     None, 'current', ['A3.4', 'B2.7'], None, 'Project knowledge (PSCP_Guidebook_Final_Version.pdf)',
     'Method for keeping essential City functions running during disruption.'),
    ('eoc-checklist', 'EOC Capability Assessment Checklist', 'framework', None, None, None, None, 'current',
     ['B2.1', 'B2.3', 'B2.4', 'B7.6'], None, 'reference/EOC Capability Assessment Checklist.docx',
     'Checklist used for the Part B facility and capability review.'),
    ('iat-workbook', 'Institutional Assessment Tool scoring workbook, IAT v1.1 (file v1-3), October 2026 baseline', 'report',
     None, 'IAT v1.1', 'CDRRMO / MOP consultant', '2026-10-30', 'current', [], None,
     'reference/Ormoc_CDRRMO-EOC_Institutional_Assessment_Tool_v1-3.xlsx',
     '140 indicators with the October 2026 baseline scores, evidence, gaps and action plan. Loaded as the "2026 Baseline" cycle.'),
    ('iat-sheet-v10', 'Institutional Assessment Tool scoring workbook v1.0 (Google Sheet, 108 indicators)', 'report', None,
     'IAT v1.0', 'MOP consultant', None, 'superseded', [],
     'https://docs.google.com/spreadsheets/d/1VLXdaFpgkzIlYwxhcb_-tM9XxgkR7qkr-qbQ7jBVg_U', None,
     'Superseded by IAT v1.1 (140 indicators, adds RA 12287 IRR requirements and Part C interface indicators).'),
    ('iat-methodology', 'IAT Standards, Policy Basis, Assessment Framework and Methodology (v1.0 draft)', 'report', None,
     'v1.0', 'MOP consultant, with CDRRMO', None, 'draft', [],
     'https://docs.google.com/document/d/189Eojg3kZQwe71Kq542K3rfNAyvOioLlF0nYCM6tgz0', None,
     'Methodology for the IAT. Still describes v1.0 (108 indicators); to be updated to v1.1 (140 indicators, Part C).'),
    ('mrt', 'Master Requirements Traceability and 140-Indicator Data Dictionary v1.1', 'report', None, 'v1.1',
     'MOP consultant', None, 'current', [], None,
     'reference/Ormoc_CDRRMO_EOC_Master_Requirements_Traceability_and_140_Indicator_Data_Dictionary_v1.1.xlsx',
     'MR-01 to MR-18, indicator-to-requirement map and data dictionary.'),
    ('inception-report', 'Inception Report: MOO for the EOC and CDRRMO of Ormoc City (Deliverable 1)', 'report', None,
     'Deliverable 1 of 5', 'MOP consultant', '2026-10-09', 'current', [],
     'https://docs.google.com/document/d/1ZNS_M5CAu51iEr_tDi9o6lJ80KWjkqh2f-kb-EDrovM', None,
     'Methodology, work plan and MOO outline for the 14-week engagement.'),
    ('work-plan', 'Project implementation plan (Gantt chart) and contract schedule', 'record', None, None,
     'MOP consultant', '2026-10-30', 'current', [],
     'https://docs.google.com/spreadsheets/d/1MrevK4TH0ECE1VGqvKRIx6qxD3CwiIfXVnsK3iOVXAQ', None,
     'Phases, tasks, milestones and progress billing.'),
    ('moo-ch1-v09', 'MOO Chapter 1: Legal mandate, DRRM governance and institutional accountability, with Implementation Package (v0.9 consolidated draft)',
     'report', None, 'v0.9', 'MOP consultant', '2026-10-31', 'draft', ['A1.1', 'A1.2', 'A1.3', 'A1.4', 'A1.6', 'A1.8', 'B1.1', 'B6.1'],
     None, 'reference/Ormoc_CDRRMO_EOC_MOO_Chapter1_Consolidated_v0.9.docx',
     'Draft chapter and annexes 1-A to 1-G. Loaded into the Library as draft versions for review.'),
    ('impl-package-long', 'Institutional Implementation Package for Chapter 1, long-form drafts (ordinance, EOs, SOPs, RACI, templates)',
     'report', None, None, 'MOP consultant', '2026-10-31', 'draft', ['A1.1', 'A1.2', 'A1.6', 'A1.8', 'B1.1'], None,
     'reference/moo/Implementation_Package_long_form.md',
     'Full-length version of the seven instruments. The consolidated v0.9 annexes are the shorter controlled set.'),
]


def vault_sql(known_ids):
    vals = []
    for key, title, cat, ref, num, issuer, date, status, inds, url, loc, summary in VAULT:
        bad = [i for i in inds if i not in known_ids]
        assert not bad, (key, bad)
        vals.append(f"({lit(uid('vault', key))}::uuid, {lit(title)}, {lit(cat)}, {lit(ref)}, {lit(num)}, {lit(issuer)}, "
                    f"{lit(date)}::date, {lit(summary)}, {lit(inds) if inds else 'array[]::text[]'}, {lit(url)}, {lit(loc)}, {lit(status)})")
    return '\n'.join([
        '-- GENERATED by scripts/generate_operational_seed.py. One-time Document Vault register of the engagement\'s',
        '-- source documents. Safe to run again: existing entries are left alone. Attach the files in the app.',
        'begin;',
        'insert into public.vault_documents (id, title, category, ref_key, doc_number, issuer, doc_date, summary, indicator_ids, url, location, status) values',
        ',\n'.join(vals), 'on conflict (id) do nothing;', 'commit;', '',
        f'select count(*) as vault_documents from public.vault_documents;  -- at least {len(VAULT)}', ''])


def main():
    ref = json.loads((ROOT / 'src' / 'data' / 'reference.json').read_text())
    known = [i['id'] for i in ref['indicators']]
    rows, plan, profile = read_workbook()
    OUT.mkdir(parents=True, exist_ok=True)
    sql, held, n_act, n_ev = baseline_sql(rows, plan, profile, known)
    for old in OUT.glob('*.sql'):
        old.unlink()
    names = ['01_iat_baseline_responses.sql', '02_iat_baseline_evidence.sql', '03_iat_baseline_scores_actions.sql']
    pieces = sql.split(PART)
    assert len(pieces) == 3
    head = '\n'.join(sql.split('\n')[:4]) + '\n'
    for k, (name, body) in enumerate(zip(names, pieces), 1):
        text = (body if k == 1 else head + f'-- Part {k} of 3. Run the parts in order.\n' + body.lstrip('\n'))
        if k == 1:
            text = text.replace('begin;', '-- Part 1 of 3. Run the parts in order.\nbegin;', 1)
        assert len(text.encode()) < 100_000, (name, len(text.encode()))
        (OUT / name).write_text(text)
    (OUT / '04_vault_register.sql').write_text(vault_sql(set(known)))
    print(f'baseline: {len(known)} responses, {n_ev} evidence notes, {n_act} actions; '
          f'{len(held)} held below the workbook score')
    for h in held:
        print('  held', *h)


if __name__ == '__main__':
    main()
